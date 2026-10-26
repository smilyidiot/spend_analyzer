"""Render an analytics report to an Excel workbook or a PDF document.

Both functions take the dict produced by :func:`app.analytics.build_report`
and return ``bytes`` so callers can stream them over HTTP or write to disk.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_table(ws, headers: list[str], rows: list[list], start_row: int = 1) -> int:
    """Write a header + rows block, style the header, return the next free row."""
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        width = max([len(str(headers[col - 1]))] + [len(str(r[col - 1])) for r in rows] or [0])
        ws.column_dimensions[letter].width = min(max(width + 2, 12), 40)
    return start_row + len(rows) + 2


def to_excel(report: dict) -> bytes:
    """Build a multi-sheet Excel workbook with a category bar chart."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Spend Analyzer — Summary"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    summ = report["summary"]
    rows = [
        ["Total spend", f"{summ['currency']} {summ['total_spend']:,.2f}"],
        ["Transactions", summ["transaction_count"]],
        ["Average transaction", f"{summ['currency']} {summ['average_transaction']:,.2f}"],
        ["Distinct vendors", summ["distinct_vendors"]],
        ["Distinct categories", summ["distinct_categories"]],
        ["Period", f"{summ['start_date']} to {summ['end_date']}"],
        ["Potential recovery", f"{summ['currency']} {report['potential_recovery']:,.2f}"],
    ]
    _write_table(ws, ["Metric", "Value"], rows, start_row=3)
    for i, note in enumerate(report["insights"], start=12):
        ws.cell(row=i, column=1, value=f"• {note}")

    cat_ws = wb.create_sheet("By Category")
    cat_rows = [[r["category"], r["total"], r["count"], r["share_pct"]] for r in report["by_category"]]
    next_row = _write_table(cat_ws, ["Category", "Total", "Count", "Share %"], cat_rows)
    if cat_rows:
        chart = BarChart()
        chart.title = "Spend by category"
        chart.type = "bar"
        chart.legend = None
        data = Reference(cat_ws, min_col=2, min_row=1, max_row=len(cat_rows) + 1)
        cats = Reference(cat_ws, min_col=1, min_row=2, max_row=len(cat_rows) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height, chart.width = 8, 16
        cat_ws.add_chart(chart, f"A{next_row}")

    ven_ws = wb.create_sheet("By Vendor")
    _write_table(
        ven_ws,
        ["Vendor", "Total", "Count"],
        [[r["vendor"], r["total"], r["count"]] for r in report["by_vendor"]],
    )

    mon_ws = wb.create_sheet("By Month")
    _write_table(
        mon_ws,
        ["Month", "Total", "MoM change %"],
        [[r["month"], r["total"], r["mom_change_pct"]] for r in report["by_month"]],
    )

    if report["duplicate_suspects"]:
        dup_ws = wb.create_sheet("Duplicates")
        _write_table(
            dup_ws,
            ["Vendor", "Amount", "Occurrences", "Dates", "Recoverable"],
            [
                [r["vendor"], r["amount"], r["occurrences"], ", ".join(r["dates"]), r["recoverable"]]
                for r in report["duplicate_suspects"]
            ],
        )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_pdf(report: dict) -> bytes:
    """Build a one-to-two page PDF summary report."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title="Spend Analyzer Report")
    styles = getSampleStyleSheet()
    story = []

    summ = report["summary"]
    story.append(Paragraph("Spend Analyzer — Report", styles["Title"]))
    story.append(Paragraph(f"Period: {summ['start_date']} to {summ['end_date']}", styles["Normal"]))
    story.append(Spacer(1, 8 * mm))

    kpi = [
        ["Total spend", f"{summ['currency']} {summ['total_spend']:,.0f}"],
        ["Transactions", str(summ["transaction_count"])],
        ["Average transaction", f"{summ['currency']} {summ['average_transaction']:,.0f}"],
        ["Potential recovery", f"{summ['currency']} {report['potential_recovery']:,.0f}"],
    ]
    story.append(_pdf_table(["Metric", "Value"], kpi))
    story.append(Spacer(1, 6 * mm))

    story.append(Paragraph("Spend by category", styles["Heading2"]))
    cat_rows = [[r["category"], f"{r['total']:,.0f}", f"{r['share_pct']:.0f}%"] for r in report["by_category"]]
    story.append(_pdf_table(["Category", "Total", "Share"], cat_rows))
    story.append(Spacer(1, 6 * mm))

    if report["insights"]:
        story.append(Paragraph("Insights", styles["Heading2"]))
        for note in report["insights"]:
            story.append(Paragraph(f"• {note}", styles["Normal"]))

    doc.build(story)
    return buffer.getvalue()


def _pdf_table(headers: list[str], rows: list[list[str]]) -> Table:
    table = Table([headers] + rows, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9CED6")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F4F7")]),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("PADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table
